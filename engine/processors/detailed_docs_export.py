"""
The narrative companion to the FINAL_PACKAGE spreadsheet -- a proper
readable Word document explaining each strategy the way it was actually
taught, not condensed into spreadsheet cells. Reuses the exact same spec
content already captured per strategy (purpose, entry_conditions,
examples, open_questions, etc.) -- this is a different PRESENTATION of
that same evidence-backed content, not a rewrite from scratch, so it
can never drift out of sync with the structured data everything else in
this package is built from.
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from processors.final_package_export import CONCEPT_GROUPS, CONCEPT_BLURBS, _group_for, _load_specs

HEADING_COLOR = RGBColor(0x1F, 0x29, 0x37)


def _add_field(doc, label, value):
    if not value:
        return
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    p.add_run(str(value))


def _add_structured_dependency(doc, label, dep: dict):
    if not dep or not dep.get("required"):
        return
    p = doc.add_paragraph()
    run = p.add_run(f"{label} dependency: ")
    run.bold = True
    bits = []
    for key in ("workbook", "sheet", "name", "platform_or_source", "formula_output", "output_interpretation",
                "calculation_logic"):
        val = dep.get(key)
        if val:
            bits.append(f"{key.replace('_', ' ')} = {val}")
    p.add_run("; ".join(bits) if bits else "required (details not fully specified)")
    for item in dep.get("required_inputs", []) or []:
        if isinstance(item, dict):
            doc.add_paragraph(f"  - needs: {item.get('name', '')} (source: {item.get('source', '')})", style=None)
        else:
            doc.add_paragraph(f"  - needs: {item}", style=None)


def _write_strategy_chapter(doc, s, concept_groups=None):
    spec = s["spec"]
    doc.add_heading(s["name"], level=1)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Concept group: {_group_for(s['name'], concept_groups)}   |   Confidence: {s['confidence']}   |   "
        f"Approved: {'Yes' if s['approved'] else 'No'}   |   Version: {s['version']}   |   "
        f"Source citations: {s['citation_count']}"
    )
    meta_run.italic = True
    meta_run.font.size = Pt(9)

    doc.add_heading("What This Strategy Is", level=2)
    doc.add_paragraph(spec.get("purpose", "(not recorded)"))

    doc.add_heading("Instrument & Timeframe", level=2)
    _add_field(doc, "Instrument", spec.get("instrument"))
    _add_field(doc, "Timeframe", spec.get("timeframe"))

    doc.add_heading("How It Works (Entry)", level=2)
    doc.add_paragraph(spec.get("entry_conditions", "(not recorded)"))

    doc.add_heading("Exit Rule", level=2)
    doc.add_paragraph(spec.get("exit_conditions") or "(not separately specified -- see stop-loss/target)")

    doc.add_heading("Stop-Loss", level=2)
    doc.add_paragraph(spec.get("stop_loss", "(not recorded)"))

    doc.add_heading("Target", level=2)
    doc.add_paragraph(spec.get("target", "(not recorded)"))

    if spec.get("no_trade_conditions"):
        doc.add_heading("When NOT To Trade This", level=2)
        doc.add_paragraph(spec["no_trade_conditions"])

    excel_dep = spec.get("excel_dependency") or {}
    indicator_dep = spec.get("indicator_dependency") or {}
    if excel_dep.get("required") or indicator_dep.get("required"):
        doc.add_heading("What You Need To Run This", level=2)
        _add_structured_dependency(doc, "Excel", excel_dep)
        _add_structured_dependency(doc, "Indicator", indicator_dep)

    external_refs = spec.get("external_reference_dependency") or []
    if external_refs:
        doc.add_heading("External References The Trainer Pointed To", level=2)
        for ref in external_refs:
            doc.add_paragraph(
                f"{ref.get('reference', '')} -- {ref.get('purpose', '')} "
                f"({'essential' if ref.get('essential') else 'supplementary'})",
                style="List Bullet",
            )

    examples = spec.get("examples") or []
    if examples:
        doc.add_heading("Worked Examples From The Session", level=2)
        for ex in examples:
            doc.add_paragraph(ex, style="List Bullet")

    open_qs = spec.get("open_questions") or []
    if open_qs:
        doc.add_heading("What's Still Unclear (Verify Before Trading Live)", level=2)
        for q in open_qs:
            doc.add_paragraph(q, style="List Bullet")

    citations = spec.get("source_citations") or []
    if citations:
        doc.add_heading("Source Citations", level=2)
        for c in citations:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{c.get('source_reference', '')}: ").bold = True
            p.add_run(c.get("evidence_summary", ""))

    doc.add_page_break()


def export_detailed_explanations(conn, course_id: str, output_path: str, concept_groups=None, concept_blurbs=None) -> dict:
    concept_groups = concept_groups if concept_groups is not None else CONCEPT_GROUPS
    concept_blurbs = concept_blurbs if concept_blurbs is not None else CONCEPT_BLURBS
    strategies = _load_specs(conn, course_id)

    doc = Document()
    title = doc.add_heading(f"Detailed Strategy Explanations -- {course_id}", level=0)
    subtitle = doc.add_paragraph()
    subtitle.add_run(
        "The full narrative behind each strategy in the companion spreadsheet (Main_Course_FINAL_PACKAGE.xlsx) "
        "-- how it was actually explained in the session, including worked examples and everything still "
        "flagged as needing your own verification. Every claim here traces back to an exact transcript "
        "timestamp or document location; see the per-strategy JSON files in 04_STRATEGY_RECONSTRUCTION/ for "
        "the complete evidence record."
    ).italic = True
    doc.add_page_break()

    by_group = {}
    for s in strategies:
        by_group.setdefault(_group_for(s["name"], concept_groups), []).append(s)

    for group in list(concept_groups.keys()) + [g for g in by_group if g not in concept_groups]:
        members = by_group.get(group)
        if not members:
            continue
        doc.add_heading(group, level=1)
        blurb = concept_blurbs.get(group)
        if blurb:
            p = doc.add_paragraph()
            p.add_run(blurb).italic = True
        doc.add_page_break()
        for s in members:
            _write_strategy_chapter(doc, s, concept_groups)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    return {
        "output_path": output_path,
        "total_strategies": len(strategies),
    }
