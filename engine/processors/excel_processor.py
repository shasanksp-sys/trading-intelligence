"""
Excel processor.

Per the architecture docs, Excel needs SPECIAL care: some sheets are just
running notes, others encode real trading logic (e.g. TTS calculations).
This processor preserves BOTH the formula and the calculated value for
every cell, plus embedded images -- never collapses a formula down to just
its result, because the formula IS the evidence of the logic.

Legacy .xls (pre-2007 binary format) is classified as file_type='excel'
the same as .xlsx/.xlsm at ingestion. openpyxl -- the library used for
.xlsx/.xlsm -- cannot open the old binary format at all, but xlrd 2.x
(despite dropping .xlsx support in that same version) still reads .xls
just fine -- confirmed directly against a real course file. A .xls sheet
is read via _process_legacy_xls below (values only -- BIFF's compiled
formula bytecode isn't practical to decompile back to a readable formula
string the way openpyxl exposes it for .xlsx, so only calculated values
are captured for this path, not formula source).
"""

import uuid
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import xlrd

from processors.text_export import safe_filename, write_text_export
from processors.image_processor import ocr_image


def _process_legacy_xls(conn, file_id: str, path: str, output_dir: str = None) -> dict:
    book = xlrd.open_workbook(path)
    extractions = []

    for sheet in book.sheets():
        for row_idx in range(sheet.nrows):
            for col_idx in range(sheet.ncols):
                cell = sheet.cell(row_idx, col_idx)
                if cell.value in (None, ""):
                    continue
                coord = f"{xlrd.formula.colname(col_idx)}{row_idx + 1}"
                location_ref = f"{sheet.name}!{coord}"
                extractions.append(
                    (str(uuid.uuid4()), file_id, location_ref, "value", f"VALUE: {cell.value}", 1.0)
                )

    conn.executemany(
        """INSERT INTO document_extractions
           (extraction_id, file_id, location_ref, extraction_type, content, confidence)
           VALUES (?, ?, ?, ?, ?, ?)""",
        extractions,
    )
    conn.execute(
        "UPDATE files SET status = 'SUCCESS', stage_completed = 'excel_extracted', processed_at = ? WHERE file_id = ?",
        (datetime.now(timezone.utc).isoformat(), file_id),
    )
    conn.commit()

    if output_dir:
        original_name = conn.execute("SELECT original_name FROM files WHERE file_id = ?", (file_id,)).fetchone()[0]
        lines = [f"Excel extraction (legacy .xls) -- {original_name}", "=" * 60, ""]
        for loc_ref, extraction_type, content, _confidence in [(e[2], e[3], e[4], e[5]) for e in extractions]:
            lines.append(f"[{loc_ref}] ({extraction_type})\n{content}\n")
        export_path = write_text_export(output_dir, f"{safe_filename(original_name)}.txt", "\n".join(lines))
        print(f"  [export] readable Excel summary written -> {export_path}")

    return {
        "sheets_processed": book.nsheets,
        "cells_extracted": len(extractions),
        "embedded_images_ocrd": 0,
    }


def process_excel_file(conn, file_id: str, path: str, output_dir: str = None) -> dict:
    """
    Extracts every non-empty cell (formula + value) from every sheet, plus
    embedded images -- which are now actually OCR'd (previously they were
    only flagged with a "route to OCR" placeholder that nothing ever
    picked up, so any text in a pasted chart/screenshot was silently
    invisible to reconstruction). Writes rows to document_extractions.
    Returns a summary dict.
    """
    ext = Path(path).suffix.lower()
    if ext == ".xls":
        return _process_legacy_xls(conn, file_id, path, output_dir)

    # Load twice: once for formulas (data_only=False), once for calculated
    # values (data_only=True) -- openpyxl can't give both from one load.
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    wb_values = openpyxl.load_workbook(path, data_only=True)

    extractions = []
    image_count = 0

    for sheet_name in wb_formulas.sheetnames:
        ws_f = wb_formulas[sheet_name]
        ws_v = wb_values[sheet_name]

        for row in ws_f.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value_cell = ws_v[cell.coordinate]
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")

                location_ref = f"{sheet_name}!{cell.coordinate}"
                if is_formula:
                    content = f"FORMULA: {cell.value}  |  CALCULATED_VALUE: {value_cell.value}"
                    extraction_type = "formula"
                else:
                    content = f"VALUE: {cell.value}"
                    extraction_type = "value"

                extractions.append(
                    (str(uuid.uuid4()), file_id, location_ref, extraction_type, content, 1.0)
                )

        # embedded images (charts/screenshots pasted into the sheet) --
        # actually extracted and OCR'd now, not just flagged. Previously
        # this only wrote a "route to OCR" placeholder that nothing ever
        # picked up, so text in a pasted screenshot (e.g. a chart with a
        # stop-loss level marked on it) was invisible to reconstruction.
        image_dir = Path(output_dir) / "embedded_images" if output_dir else None
        if image_dir:
            image_dir.mkdir(parents=True, exist_ok=True)
        for image in getattr(ws_f, "_images", []):
            image_count += 1
            location_ref = f"{sheet_name}!embedded_image_{image_count}"
            try:
                image_bytes = image._data()
                ext = Path(image.path).suffix or ".png"
                if image_dir:
                    img_path = image_dir / f"{safe_filename(sheet_name)}_img{image_count}{ext}"
                    img_path.write_bytes(image_bytes)
                    ocr_text = ocr_image(str(img_path))
                else:
                    # no output_dir given (e.g. a direct/test call) -- OCR via a temp file
                    import tempfile
                    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                        tmp.write(image_bytes)
                        tmp_path = tmp.name
                    ocr_text = ocr_image(tmp_path)
                    Path(tmp_path).unlink(missing_ok=True)
                content = ocr_text if ocr_text.strip() else "(embedded image found, but no text detected by OCR)"
            except Exception as e:
                content = f"(embedded image found in sheet '{sheet_name}', but OCR failed: {e})"
            extractions.append(
                (str(uuid.uuid4()), file_id, location_ref, "embedded_image", content, 1.0)
            )

    conn.executemany(
        """INSERT INTO document_extractions
           (extraction_id, file_id, location_ref, extraction_type, content, confidence)
           VALUES (?, ?, ?, ?, ?, ?)""",
        extractions,
    )
    conn.execute(
        "UPDATE files SET status = 'SUCCESS', stage_completed = 'excel_extracted', processed_at = ? WHERE file_id = ?",
        (datetime.now(timezone.utc).isoformat(), file_id),
    )
    conn.commit()

    if output_dir:
        original_name = conn.execute("SELECT original_name FROM files WHERE file_id = ?", (file_id,)).fetchone()[0]
        lines = [f"Excel extraction -- {original_name}", "=" * 60, ""]
        for loc_ref, extraction_type, content, _confidence in [(e[2], e[3], e[4], e[5]) for e in extractions]:
            lines.append(f"[{loc_ref}] ({extraction_type})\n{content}\n")
        export_path = write_text_export(output_dir, f"{safe_filename(original_name)}.txt", "\n".join(lines))
        print(f"  [export] readable Excel summary written -> {export_path}")

    return {
        "sheets_processed": len(wb_formulas.sheetnames),
        "cells_extracted": len(extractions) - image_count,
        "embedded_images_ocrd": image_count,
    }


if __name__ == "__main__":
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from db.schema import get_connection

    conn = get_connection("TradingIntelligence.db")
    row = conn.execute(
        "SELECT file_id, original_path FROM files WHERE file_type = 'excel' LIMIT 1"
    ).fetchone()
    file_id, path = row
    result = process_excel_file(conn, file_id, path)
    print("Excel processing result:", result)

    print("\nSample extracted rows:")
    for r in conn.execute(
        "SELECT location_ref, extraction_type, content FROM document_extractions WHERE file_id = ?",
        (file_id,),
    ):
        print(" ", r)
    conn.close()
