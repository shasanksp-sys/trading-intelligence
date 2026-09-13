"""
PPTX processor.

Extracts text from every slide (titles, body text, text boxes) plus
speaker notes, each tagged with its slide number for traceability --
same principle as the PDF processor's page numbers and the Excel
processor's cell references.

Also extracts and OCRs every picture shape on a slide (charts,
screenshots, or marked-up diagrams pasted directly onto a slide as an
image rather than typed as text) -- previously these were skipped
entirely, the same class of gap that was found and fixed for embedded
images in Excel and PDF. Picture shapes are matched recursively inside
grouped shapes too, since trainers commonly group an image with a text
label on top of it.
"""

import io
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from processors.text_export import safe_filename, write_text_export
from processors.image_processor import ocr_image

# .ppsx ("PowerPoint Show") is byte-for-byte the same OOXML package
# structure as .pptx -- same slide/shape/table parts -- except
# [Content_Types].xml declares the main part as
# ".presentationml.slideshow.main+xml" instead of
# ".presentationml.presentation.main+xml". python-pptx's Presentation()
# checks that content type and refuses to open it, even though nothing
# about the actual slide content differs. Rather than skip every .ppsx
# file in a course (found via 19 stuck files in the Sameer Dharaskar
# 2023 course alone, 42 more in the 2022 course), patch just that one
# string in-memory and hand python-pptx a corrected copy.
_SLIDESHOW_CT = b"presentationml.slideshow.main+xml"
_PRESENTATION_CT = b"presentationml.presentation.main+xml"


def _open_presentation(path: str):
    if not str(path).lower().endswith(".ppsx"):
        return Presentation(path)

    with open(path, "rb") as f:
        original_bytes = f.read()
    src = zipfile.ZipFile(io.BytesIO(original_bytes))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_SLIDESHOW_CT, _PRESENTATION_CT)
            dst.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def _iter_picture_shapes(shapes):
    """Yields every picture shape, descending into grouped shapes recursively."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_picture_shapes(shape.shapes)


def process_pptx_file(conn, file_id: str, path: str, output_dir: str = None) -> dict:
    prs = _open_presentation(path)
    extractions = []
    image_count = 0

    image_dir = Path(output_dir) / "embedded_images" if output_dir else None
    if image_dir:
        image_dir.mkdir(parents=True, exist_ok=True)

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_texts.append(text)
            if shape.has_table:
                table_lines = []
                for row in shape.table.rows:
                    table_lines.append(" | ".join(cell.text.strip() for cell in row.cells))
                if table_lines:
                    slide_texts.append("\n".join(table_lines))

        if slide_texts:
            extractions.append(
                (str(uuid.uuid4()), file_id, f"slide_{slide_num}", "text",
                 "\n".join(slide_texts), 1.0)
            )

        # speaker notes -- trainers sometimes put the actual explanation here
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                extractions.append(
                    (str(uuid.uuid4()), file_id, f"slide_{slide_num}_notes", "text",
                     notes_text, 1.0)
                )

        # embedded pictures (charts/screenshots pasted directly onto the slide) --
        # extracted and OCR'd, not just skipped. Same fix already applied to
        # Excel and PDF: a placeholder with no real content is as good as
        # invisible to strategy reconstruction, so this reads the actual text.
        for pic_idx, picture in enumerate(_iter_picture_shapes(slide.shapes), start=1):
            image_count += 1
            location_ref = f"slide_{slide_num}!embedded_image_{pic_idx}"
            try:
                image = picture.image
                ext = image.ext or "png"
                if image_dir:
                    img_path = image_dir / f"slide{slide_num}_img{pic_idx}.{ext}"
                    img_path.write_bytes(image.blob)
                    ocr_text = ocr_image(str(img_path))
                else:
                    import tempfile
                    with tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False) as tmp:
                        tmp.write(image.blob)
                        tmp_path = tmp.name
                    ocr_text = ocr_image(tmp_path)
                    Path(tmp_path).unlink(missing_ok=True)
                content = ocr_text if ocr_text.strip() else "(embedded image found, but no text detected by OCR)"
            except Exception as e:
                content = f"(embedded image found on slide {slide_num}, but OCR failed: {e})"
            extractions.append(
                (str(uuid.uuid4()), file_id, location_ref, "embedded_image", content, 1.0)
            )

    conn.executemany(
        """INSERT INTO document_extractions
           (extraction_id, file_id, location_ref, extraction_type, content, confidence)
           VALUES (?, ?, ?, ?, ?, ?)""",
        extractions,
    )
    conn.execute(
        "UPDATE files SET status = 'SUCCESS', stage_completed = 'pptx_extracted', processed_at = ? WHERE file_id = ?",
        (datetime.now(timezone.utc).isoformat(), file_id),
    )
    conn.commit()

    if output_dir:
        original_name = conn.execute("SELECT original_name FROM files WHERE file_id = ?", (file_id,)).fetchone()[0]
        lines = [f"PowerPoint extraction -- {original_name}", "=" * 60, ""]
        for e in extractions:
            lines.append(f"[{e[2]}]\n{e[4]}\n")
        export_path = write_text_export(output_dir, f"{safe_filename(original_name)}.txt", "\n".join(lines))
        print(f"  [export] readable PPTX summary written -> {export_path}")

    return {
        "slides_processed": len(prs.slides),
        "text_blocks_extracted": len(extractions) - image_count,
        "embedded_images_ocrd": image_count,
    }


if __name__ == "__main__":
    import sys
    from pathlib import Path
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from db.schema import get_connection

    conn = get_connection("TradingIntelligence.db")
    row = conn.execute("SELECT file_id, original_path FROM files WHERE file_type = 'pptx' LIMIT 1").fetchone()
    if row:
        file_id, path = row
        result = process_pptx_file(conn, file_id, path)
        print("PPTX processing result:", result)
    else:
        print("No pptx file found in database to test with.")
    conn.close()
